import openpyxl
from pathlib import Path
from typing import List, Dict, Optional
from logger import logger


class ExcelHandler:
    def __init__(self, config):
        self.config = config
        self.workbook = None
        self.sheet = None
        self.products = []

    def load_products(self):
        self.load_products_from(self.config.input_file)

    def load_products_from(self, file_path: Path):
        logger.info(f"Loading products from {file_path}")
        self.workbook = openpyxl.load_workbook(file_path)
        self.sheet = self.workbook.active
        self.products = []

        # قراءة الصفوف (يفترض وجود العناوين في الصف الأول)
        for idx, row in enumerate(self.sheet.iter_rows(min_row=2, values_only=True), start=2):
            product = {
                'index': idx,
                'name': str(row[1]) if len(row) > 1 and row[1] is not None else '',
                'barcode': str(row[0]) if len(row) > 0 and row[0] is not None else '',
                'image_url': row[4] if len(row) > 4 and row[4] is not None else '',
                'status': '',
                'confidence': 0,
                'download_path': '',
            }
            self.products.append(product)

    def get_products_without_images(self) -> List[Dict]:
        return [
            p for p in self.products
            if not p.get('image_url') or p.get('image_url') in ['', 'Not Found', None]
        ]

    def update_image_url(self, index: int, url: str):
        # العمود الخامس (E) هو link_photo في ملف الإكسل
        self.sheet.cell(row=index, column=5, value=url)

    def update_image_status(self, index: int, status: str):
        pass

    def update_confidence(self, index: int, confidence: float):
        pass

    def update_download_path(self, index: int, path: str):
        pass

    def save_products(self):
        self.workbook.save(self.config.output_file)
        logger.info(f"Saved results to {self.config.output_file}")

    def save_checkpoint(self):
        self.workbook.save(self.config.output_file)

    def get_statistics(self):
        total = len(self.products)
        with_images = sum(
            1 for p in self.products
            if p.get('image_url') and p.get('image_url') not in ['', 'Not Found', None]
        )
        return {
            'total': total,
            'with_images': with_images,
            'without_images': total - with_images,
            'progress': (with_images / total * 100) if total > 0 else 0
        }